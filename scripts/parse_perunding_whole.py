# Parses the authoritative "Perunding Whole.xlsx" export from Wasiyyah into
# scripts/roster_import_v2.json, in the same record shape import-roster.mjs
# already expects (rosterId/daieId/name/rank/unitId/uplineRosterId/
# lineagePath/status/isLdpMember/flags/sourceRow), plus two new fields this
# source makes possible: email and dateLicensed (from Tempoh Mula).
#
# Deliberately does NOT read IC, Tarikh Lahir, No Tel(2), Bank, No Account —
# data minimization: only what the app's schema actually uses.
#
# isLdpMember is not present in this file at all (it was inferred from
# manual orange color-coding in the old hierarchy sheet) — carried forward
# from scripts/roster_import.json by matching daieId, since that's the only
# place this signal exists.
#
# Usage: python scripts/parse_perunding_whole.py "<path to Perunding Whole.xlsx>"

import sys
import re
import json
from collections import Counter
from datetime import datetime

import openpyxl

UNIT_BY_KUMPULAN = {
    "KAUSAR WEALTH MANAGEMENT": "kausar-wealth",
    "KAUSAR GLOBAL": "kausar-global",
    "KAUSAR ASPIRE": "kausar-aspire",
    "KAUSAR INTISAR": "kausar-intisar",
    "KAUSAR NUSRAH": "kausar-nusrah",
}

RANK_BY_LEVEL = {
    "PM": "DM",
    "PPM": "DPM",
    "KPE": "KDE",
    "KDE": "KDE",
}

UPLINE_KOD_RE = re.compile(r"\((\d+)\)?\s*$")
UPLINE_LABEL_RE = re.compile(r"^(.*?)\s*\(\d+\)?\s*$")

# PIC review (2026-09-13) of the cross_unit_upline flags this parser raised
# on the first run: the sheet's own Kumpulan (unit) column is wrong for
# these 4 people — their real unit follows their actual introducer/upline,
# per usayri.kausar@gmail.com. Overriding here (rather than hand-editing the
# generated JSON) so the fix survives a re-parse of an updated source file.
UNIT_OVERRIDES = {
    "10521": "kausar-nusrah",  # Rika Binti Tahir — was Noorashiqien Ishak's (Kausar Nusrah KDE) downline before expiring; whole lineage traces back to Kausar Nusrah, not Kausar Wealth
    "86364": "kausar-aspire",  # Mohd Osmerayusman Bin Mohamad — recruited by Azura Binti Mohd Arif, Kausar Aspire
    "86834": "kausar-intisar",  # Mohd Fahmi Bin Md Miftahuddin — introducer Mohamad Nazri Daud is Kausar Intisar
    "86840": "kausar-intisar",  # Aimi Mursyidah Binti Suria — same introducer (Mohamad Nazri Daud), Kausar Intisar
}

# Also from that same review: this cross-unit upline is real, not a data
# error — Mohd Fazli Bin Omar (Kausar Wealth) is confirmed as Intan
# Yusliyana's actual introducer/sponsor even though she became Kausar
# Intisar's KDE. Flagged distinctly so it doesn't read as still-open.
CONFIRMED_CROSS_UNIT_UPLINES = {"11353"}


def main():
    if len(sys.argv) < 2:
        print("Usage: python scripts/parse_perunding_whole.py <path to xlsx>")
        sys.exit(1)

    path = sys.argv[1]
    old_roster_path = sys.argv[2] if len(sys.argv) > 2 else "scripts/roster_import.json"

    old_ldp_by_daie_id = {}
    try:
        with open(old_roster_path, encoding="utf-8") as f:
            for r in json.load(f):
                if r.get("isLdpMember"):
                    old_ldp_by_daie_id[str(r["daieId"])] = True
    except FileNotFoundError:
        print(f"(no old roster found at {old_roster_path}, isLdpMember will default to false)")

    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["Sheet1"]
    headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
    idx = {h: i for i, h in enumerate(headers)}

    raw_rows = []
    for r in range(2, ws.max_row + 1):
        row = [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
        if row[idx["Nama"]] is None:
            continue
        raw_rows.append((r, row))

    # Pass 1: assign a stable rosterId per row and index by daieId (Kod).
    kod_counts = Counter()
    for _sourceRow, row in raw_rows:
        kod = row[idx["Kod"]]
        if kod is not None:
            kod_counts[str(kod)] += 1

    records = []
    row_by_kod = {}  # daieId(str) -> record, only for unique, non-missing kods
    for sourceRow, row in raw_rows:
        kod_raw = row[idx["Kod"]]
        daie_id = str(kod_raw) if kod_raw is not None else None
        flags = []

        if daie_id is None:
            roster_id = f"row{sourceRow}"
            flags.append("missing_id")
        elif kod_counts[daie_id] > 1:
            roster_id = f"row{sourceRow}"  # keep stable but distinct per duplicate
            flags.append("duplicate_id")
        else:
            roster_id = f"k{daie_id}"

        level = row[idx["Level"]]
        rank = RANK_BY_LEVEL.get(level)
        if rank is None:
            rank = "DM"
            flags.append(f"unknown_level:{level}")

        kumpulan = row[idx["Kumpulan"]]
        unit_id = UNIT_BY_KUMPULAN.get(kumpulan)
        if unit_id is None:
            unit_id = "kausar-wealth"
            flags.append(f"unresolved_unit:{kumpulan}")
        if daie_id in UNIT_OVERRIDES:
            unit_id = UNIT_OVERRIDES[daie_id]
            flags.append("unit_corrected_by_pic")

        mula = row[idx["Tempoh Mula"]]
        date_licensed = mula.date().isoformat() if isinstance(mula, datetime) else None

        email = row[idx["Emel"]]
        email = str(email).strip() if email else None
        if not email:
            flags.append("missing_email")

        rec = {
            "rosterId": roster_id,
            "daieId": daie_id,
            "name": str(row[idx["Nama"]]).strip(),
            "email": email,
            "rank": rank,
            "unitId": unit_id,
            "dateLicensed": date_licensed,
            "isLdpMember": old_ldp_by_daie_id.get(daie_id, False),
            "flags": flags,
            "sourceRow": sourceRow,
            "_uplineRaw": row[idx["Upline"]],
        }
        records.append(rec)
        if daie_id is not None and kod_counts[daie_id] == 1:
            row_by_kod[daie_id] = rec

    # Some Upline references point to a "(kod)" that never appears as a real
    # person's Kod anywhere in the sheet — e.g. "KAUSAR WEALTH MANAGEMENT
    # SDN. BHD. (10026)", the parent company itself, which (unlike the other
    # 4 units) has no person explicitly marked KDE. Mirrors the old
    # hierarchy import's approach: synthesize one root record per such
    # phantom entity so its direct reports nest under a real node instead of
    # each becoming a stray top-level root.
    phantom_kods = {}
    for _sourceRow, row in raw_rows:
        upline_raw = row[idx["Upline"]]
        if not upline_raw:
            continue
        m = UPLINE_KOD_RE.search(str(upline_raw))
        if not m:
            continue
        up_kod = m.group(1)
        if up_kod in row_by_kod or up_kod in phantom_kods:
            continue
        label_m = UPLINE_LABEL_RE.match(str(upline_raw))
        label = label_m.group(1).strip() if label_m else str(upline_raw)
        unit_id = next(
            (uid for name, uid in UNIT_BY_KUMPULAN.items() if name in label.upper()),
            "kausar-wealth",
        )
        phantom_kods[up_kod] = {
            "rosterId": f"phantom{up_kod}",
            "daieId": up_kod,
            "name": label,
            "email": None,
            "rank": "KDE",
            "unitId": unit_id,
            "dateLicensed": None,
            "isLdpMember": old_ldp_by_daie_id.get(up_kod, False),
            "flags": ["synthetic_root"],
            "sourceRow": None,
            "uplineRosterId": None,
        }
    records.extend(phantom_kods.values())

    # Pass 2: resolve Upline -> uplineRosterId (or null if it's a root/
    # unresolvable external reference).
    unresolved_upline_examples = Counter()
    cross_unit = []
    for rec in records:
        upline_raw = rec.pop("_uplineRaw", None)
        rec.setdefault("uplineRosterId", None)
        if upline_raw:
            m = UPLINE_KOD_RE.search(str(upline_raw))
            if m:
                up_kod = m.group(1)
                up_rec = row_by_kod.get(up_kod) or phantom_kods.get(up_kod)
                if up_rec:
                    rec["uplineRosterId"] = up_rec["rosterId"]
                    if up_rec["unitId"] != rec["unitId"]:
                        if "synthetic_root" in up_rec["flags"]:
                            # Unit's own KDE reporting to the Kausar Wealth
                            # Management parent-company entity — expected,
                            # not an anomaly (see the phantom-root comment
                            # above).
                            rec["flags"].append("cross_unit_upline_structural")
                        elif rec["daieId"] in CONFIRMED_CROSS_UNIT_UPLINES:
                            rec["flags"].append("cross_unit_upline_confirmed")
                        else:
                            rec["flags"].append("cross_unit_upline")
                            cross_unit.append((rec["daieId"], rec["name"], rec["unitId"], up_kod, up_rec["name"], up_rec["unitId"]))
            else:
                rec["flags"].append("unresolved_upline")
                unresolved_upline_examples[str(upline_raw)] += 1

    # Pass 3: lineagePath by walking uplineRosterId chains.
    by_id = {r["rosterId"]: r for r in records}

    def lineage_of(rec, seen=None):
        seen = seen or set()
        if rec["rosterId"] in seen:
            rec["flags"].append("upline_cycle")
            return []
        seen.add(rec["rosterId"])
        up_id = rec["uplineRosterId"]
        if not up_id or up_id not in by_id:
            return []
        up = by_id[up_id]
        return lineage_of(up, seen) + [up_id]

    for rec in records:
        rec["lineagePath"] = lineage_of(rec)

    # Pass 4: status, computed from dateLicensed + rank using the same
    # 24/36-month rule as lib/utils.ts's contractExpiryDate (kept in sync by
    # hand — see that file for the source of truth). No date -> active
    # (mirrors isActiveStatus treating a null dateLicensed as not-yet-expired).
    def add_months(d, months):
        month = d.month - 1 + months
        year = d.year + month // 12
        month = month % 12 + 1
        day = min(d.day, [31, 29 if year % 4 == 0 and (year % 100 != 0 or year % 400 == 0) else 28,
                          31, 30, 31, 30, 31, 31, 30, 31, 30, 31][month - 1])
        return d.replace(year=year, month=month, day=day)

    today = datetime.now().date()
    active_count = 0
    expired_count = 0
    for rec in records:
        if rec["dateLicensed"] is None:
            rec["status"] = "active"
        else:
            months = 36 if rec["rank"] == "KDE" else 24
            start = datetime.fromisoformat(rec["dateLicensed"]).date()
            expiry = add_months(start, months)
            rec["status"] = "active" if expiry >= today else "expired"
        if rec["status"] == "active":
            active_count += 1
        else:
            expired_count += 1

    out_path = "scripts/roster_import_v2.json"
    with open(out_path, "w", encoding="utf-8") as f:
        json.dump(records, f, indent=1, ensure_ascii=False)

    dupes = sum(1 for r in records if "duplicate_id" in r["flags"])
    missing_id = sum(1 for r in records if "missing_id" in r["flags"])
    missing_email = sum(1 for r in records if "missing_email" in r["flags"])
    unresolved = sum(1 for r in records if "unresolved_upline" in r["flags"])
    cross = sum(1 for r in records if "cross_unit_upline" in r["flags"])
    cross_structural = sum(1 for r in records if "cross_unit_upline_structural" in r["flags"])
    cross_confirmed = sum(1 for r in records if "cross_unit_upline_confirmed" in r["flags"])
    corrected = sum(1 for r in records if "unit_corrected_by_pic" in r["flags"])

    print(f"Wrote {out_path}: {len(records)} records")
    print(f"  active={active_count} expired={expired_count}")
    print(f"  flagged: duplicate_id={dupes} missing_id={missing_id} missing_email={missing_email} "
          f"unresolved_upline={unresolved}")
    print(f"  cross_unit_upline: unresolved={cross} structural(unit-KDE->parent company)={cross_structural} "
          f"PIC-confirmed={cross_confirmed} | unit_corrected_by_pic={corrected}")
    print(f"\nUnresolved-upline external references: {dict(unresolved_upline_examples)}")
    if cross_unit:
        print(f"\nStill-unresolved cross-unit upline pairs ({len(cross_unit)}):")
        for c in cross_unit:
            print(" ", c)


if __name__ == "__main__":
    main()
